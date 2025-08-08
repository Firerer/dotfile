#!/usr/bin/env -S uv run --script
#
# /// script
# dependencies = [
#   "python-magic",
#   "PyPDF2", 
#   "python-docx",
#   "openpyxl",
#   "beautifulsoup4",
#   "chardet",
#   "pycryptodome"
# ]
# ///

"""
Production-ready file processing tool for extracting text content from various file formats.

Supports: PDF, DOCX, XLSX, HTML, TXT, RTF
Preserves document structure and formatting where possible.
"""

import argparse
import logging
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import traceback

# Import dependencies with error handling
try:
    import magic
    import PyPDF2
    from docx import Document
    from openpyxl import load_workbook
    from bs4 import BeautifulSoup
    import chardet
except ImportError as e:
    print(f"Missing dependency: {e}")
    print("Install with: uv add python-magic PyPDF2 python-docx openpyxl beautifulsoup4 chardet")
    sys.exit(1)


def setup_logging() -> logging.Logger:
    """Configure logging with appropriate format and level."""
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s',
        datefmt='%Y-%m-%d %H:%M:%S'
    )
    return logging.getLogger(__name__)


def detect_file_type(file_path: Path) -> str:
    """
    Detect file type using python-magic.
    
    Args:
        file_path: Path to the file
        
    Returns:
        MIME type string
    """
    try:
        mime = magic.Magic(mime=True)
        return mime.from_file(str(file_path))
    except Exception:
        # Fallback to extension-based detection
        suffix = file_path.suffix.lower()
        mime_map = {
            '.pdf': 'application/pdf',
            '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            '.html': 'text/html',
            '.htm': 'text/html',
            '.txt': 'text/plain',
            '.rtf': 'application/rtf'
        }
        return mime_map.get(suffix, 'application/octet-stream')


def detect_encoding(file_path: Path) -> str:
    """
    Detect file encoding using chardet.
    
    Args:
        file_path: Path to the file
        
    Returns:
        Detected encoding string, defaults to 'utf-8'
    """
    try:
        with open(file_path, 'rb') as f:
            raw_data = f.read(10000)  # Read first 10KB for detection
            result = chardet.detect(raw_data)
            return result.get('encoding', 'utf-8') or 'utf-8'
    except Exception:
        return 'utf-8'


def clean_text(text: str) -> str:
    """
    Clean and normalize extracted text while preserving structure.
    
    Args:
        text: Raw text content
        
    Returns:
        Cleaned text with normalized whitespace
    """
    if not text:
        return ""
    
    # Normalize line breaks and remove excessive whitespace
    lines = text.split('\n')
    cleaned_lines = []
    
    for line in lines:
        cleaned_line = ' '.join(line.split())  # Normalize internal whitespace
        cleaned_lines.append(cleaned_line)
    
    # Remove excessive empty lines while preserving paragraph breaks
    result_lines = []
    prev_empty = False
    
    for line in cleaned_lines:
        if line.strip():
            result_lines.append(line)
            prev_empty = False
        elif not prev_empty:
            result_lines.append("")
            prev_empty = True
    
    return '\n'.join(result_lines).strip()


def extract_pdf_text(file_path: Path) -> str:
    """
    Extract text from PDF files using PyPDF2.
    
    Args:
        file_path: Path to PDF file
        
    Returns:
        Extracted text content
        
    Raises:
        Exception: If PDF processing fails
    """
    text_content = []
    
    with open(file_path, 'rb') as file:
        pdf_reader = PyPDF2.PdfReader(file)
        
        for page_num, page in enumerate(pdf_reader.pages, 1):
            try:
                page_text = page.extract_text()
                if page_text.strip():
                    text_content.append(f"--- Page {page_num} ---\n{page_text}")
            except Exception as e:
                logging.warning(f"Failed to extract text from page {page_num}: {e}")
                continue
    
    return '\n\n'.join(text_content)


def extract_docx_text(file_path: Path) -> str:
    """
    Extract text from DOCX files preserving structure.
    
    Args:
        file_path: Path to DOCX file
        
    Returns:
        Extracted text with preserved formatting
        
    Raises:
        Exception: If DOCX processing fails
    """
    doc = Document(file_path)
    text_parts = []
    
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            # Preserve heading structure
            if paragraph.style.name.startswith('Heading'):
                level = paragraph.style.name.replace('Heading ', '')
                text_parts.append(f"{'#' * int(level) if level.isdigit() else '#'} {text}")
            else:
                text_parts.append(text)
    
    # Extract table content
    for table in doc.tables:
        table_text = []
        for row in table.rows:
            row_text = ' | '.join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                table_text.append(row_text)
        
        if table_text:
            text_parts.append("\n--- Table ---")
            text_parts.extend(table_text)
            text_parts.append("--- End Table ---\n")
    
    return '\n\n'.join(text_parts)


def extract_xlsx_text(file_path: Path) -> str:
    """
    Extract text from Excel files including all worksheets.
    
    Args:
        file_path: Path to XLSX file
        
    Returns:
        Extracted text from all worksheets
        
    Raises:
        Exception: If Excel processing fails
    """
    workbook = load_workbook(file_path, data_only=True)
    text_parts = []
    
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        text_parts.append(f"=== Worksheet: {sheet_name} ===")
        
        sheet_content = []
        for row in worksheet.iter_rows(values_only=True):
            row_text = ' | '.join(str(cell) if cell is not None else '' for cell in row)
            if row_text.strip() and row_text != ' | ' * (len(row) - 1):
                sheet_content.append(row_text)
        
        if sheet_content:
            text_parts.extend(sheet_content)
        else:
            text_parts.append("(Empty worksheet)")
        
        text_parts.append("")  # Add spacing between sheets
    
    return '\n'.join(text_parts)


def extract_html_text(file_path: Path) -> str:
    """
    Extract text from HTML files preserving structure.
    
    Args:
        file_path: Path to HTML file
        
    Returns:
        Extracted text with preserved structure
        
    Raises:
        Exception: If HTML processing fails
    """
    encoding = detect_encoding(file_path)
    
    with open(file_path, 'r', encoding=encoding) as file:
        soup = BeautifulSoup(file.read(), 'html.parser')
    
    # Remove script and style elements
    for script in soup(["script", "style"]):
        script.extract()
    
    text_parts = []
    
    # Extract title
    title = soup.find('title')
    if title:
        text_parts.append(f"# {title.get_text().strip()}")
    
    # Extract structured content
    for element in soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'div', 'li']):
        text = element.get_text().strip()
        if text:
            if element.name.startswith('h'):
                level = int(element.name[1])
                text_parts.append(f"{'#' * level} {text}")
            elif element.name == 'li':
                text_parts.append(f"• {text}")
            else:
                text_parts.append(text)
    
    return '\n\n'.join(text_parts)


def extract_text_file(file_path: Path) -> str:
    """
    Extract text from plain text files with encoding detection.
    
    Args:
        file_path: Path to text file
        
    Returns:
        File content as string
        
    Raises:
        Exception: If file reading fails
    """
    encoding = detect_encoding(file_path)
    
    with open(file_path, 'r', encoding=encoding, errors='replace') as file:
        return file.read()


def extract_rtf_text(file_path: Path) -> str:
    """
    Extract text from RTF files (basic implementation).
    
    Args:
        file_path: Path to RTF file
        
    Returns:
        Extracted text content
        
    Raises:
        Exception: If RTF processing fails
    """
    encoding = detect_encoding(file_path)
    
    with open(file_path, 'r', encoding=encoding, errors='replace') as file:
        content = file.read()
    
    # Basic RTF parsing - remove control words and groups
    import re
    
    # Remove RTF control words
    content = re.sub(r'\\[a-z]+\d*\s?', ' ', content)
    # Remove group markers
    content = re.sub(r'[{}]', ' ', content)
    # Clean up whitespace
    content = re.sub(r'\s+', ' ', content)
    
    return content.strip()


def get_extractor_function(mime_type: str) -> Optional[callable]:
    """
    Get the appropriate extraction function for a given MIME type.
    
    Args:
        mime_type: MIME type of the file
        
    Returns:
        Extraction function or None if unsupported
    """
    extractors = {
        'application/pdf': extract_pdf_text,
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document': extract_docx_text,
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': extract_xlsx_text,
        'text/html': extract_html_text,
        'text/plain': extract_text_file,
        'application/rtf': extract_rtf_text,
        'text/rtf': extract_rtf_text,
    }
    
    return extractors.get(mime_type)


def process_file(file_path: Path, logger: logging.Logger) -> bool:
    """
    Process a single file and extract its text content.
    
    Args:
        file_path: Path to the file to process
        logger: Logger instance
        
    Returns:
        True if processing succeeded, False otherwise
    """
    if not file_path.exists():
        logger.error(f"File not found: {file_path}")
        return False
    
    if not file_path.is_file():
        logger.warning(f"Skipping non-file: {file_path}")
        return False
    
    try:
        # Detect file type
        mime_type = detect_file_type(file_path)
        logger.info(f"Processing {file_path.name} (detected as {mime_type})")
        
        # Get appropriate extractor
        extractor = get_extractor_function(mime_type)
        if not extractor:
            logger.warning(f"Unsupported file type {mime_type} for {file_path.name}")
            return False
        
        # Extract text
        extracted_text = extractor(file_path)
        
        if not extracted_text.strip():
            logger.warning(f"No text content extracted from {file_path.name}")
            return False
        
        # Clean the extracted text
        cleaned_text = clean_text(extracted_text)
        
        # Generate output file path
        output_path = file_path.parent / f"{file_path.stem}_extracted.txt"
        
        # Write extracted text to file
        with open(output_path, 'w', encoding='utf-8') as output_file:
            output_file.write(cleaned_text)
        
        logger.info(f"Text extracted to: {output_path}")
        return True
        
    except Exception as e:
        logger.error(f"Failed to process {file_path.name}: {e}")
        logger.debug(f"Error details: {traceback.format_exc()}")
        return False


def main() -> int:
    """
    Main function to handle command-line arguments and process files.
    
    Returns:
        Exit code (0 for success, 1 for failure)
    """
    parser = argparse.ArgumentParser(
        description="Extract text content from various file formats",
        epilog="Supported formats: PDF, DOCX, XLSX, HTML, TXT, RTF"
    )
    parser.add_argument(
        'files',
        nargs='+',
        type=Path,
        help='File paths to process'
    )
    parser.add_argument(
        '-v', '--verbose',
        action='store_true',
        help='Enable verbose logging'
    )
    
    args = parser.parse_args()
    
    # Setup logging
    logger = setup_logging()
    if args.verbose:
        logger.setLevel(logging.DEBUG)
    
    # Process files
    total_files = len(args.files)
    successful = 0
    failed = 0
    
    logger.info(f"Starting batch processing of {total_files} file(s)")
    
    for i, file_path in enumerate(args.files, 1):
        logger.info(f"Progress: {i}/{total_files}")
        
        if process_file(file_path, logger):
            successful += 1
        else:
            failed += 1
    
    # Summary
    logger.info(f"Processing complete: {successful} successful, {failed} failed")
    
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
